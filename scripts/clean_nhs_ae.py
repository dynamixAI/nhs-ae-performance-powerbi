#!/usr/bin/env python3
"""
clean_nhs_ae.py

Combines and cleans the raw NHS England monthly A&E CSVs in data/raw/ and
writes a star-schema Excel workbook for Power BI to data/clean/.

Output: data/clean/nhs_ae_model.xlsx with three sheets (each an Excel table):
  fact_ae   - one row per organisation per month (the numbers)
  dim_org   - one row per organisation (code, name, region)
  dim_date  - one row per month (calendar + NHS financial-year attributes)

Usage: python3 scripts/clean_nhs_ae.py
"""

import re
import sys
from pathlib import Path

import pandas as pd
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.table import Table, TableStyleInfo

# ---------- Settings ----------
PROJECT = Path(__file__).resolve().parent.parent
RAW_DIR = PROJECT / "data" / "raw"
CLEAN_DIR = PROJECT / "data" / "clean"

# Original NHS header (after trimming spaces) -> clean, short name.
# Only these 22 columns are kept, which also drops the empty junk columns in Sep 2024.
COLUMN_MAP = {
    "Period": "period_raw",
    "Org Code": "org_code",
    "Parent Org": "region",
    "Org name": "org_name",
    "A&E attendances Type 1": "att_type1",
    "A&E attendances Type 2": "att_type2",
    "A&E attendances Other A&E Department": "att_other",
    "A&E attendances Booked Appointments Type 1": "att_booked_type1",
    "A&E attendances Booked Appointments Type 2": "att_booked_type2",
    "A&E attendances Booked Appointments Other Department": "att_booked_other",
    "Attendances over 4hrs Type 1": "over4h_type1",
    "Attendances over 4hrs Type 2": "over4h_type2",
    "Attendances over 4hrs Other Department": "over4h_other",
    "Attendances over 4hrs Booked Appointments Type 1": "over4h_booked_type1",
    "Attendances over 4hrs Booked Appointments Type 2": "over4h_booked_type2",
    "Attendances over 4hrs Booked Appointments Other Department": "over4h_booked_other",
    "Patients who have waited 4-12 hs from DTA to admission": "dta_wait_4_12h",
    "Patients who have waited 12+ hrs from DTA to admission": "dta_wait_12h_plus",
    "Emergency admissions via A&E - Type 1": "em_adm_type1",
    "Emergency admissions via A&E - Type 2": "em_adm_type2",
    "Emergency admissions via A&E - Other A&E department": "em_adm_other",
    "Other emergency admissions": "em_adm_other_routes",
}
TEXT_COLS = ["org_code", "region", "org_name"]
LABEL_COLS = ["period_raw"] + TEXT_COLS          # any of these may say "TOTAL"
NUMERIC_COLS = [c for c in COLUMN_MAP.values() if c not in LABEL_COLS]
ATT_COLS = [c for c in NUMERIC_COLS if c.startswith("att_")]
OVER4H_COLS = [c for c in NUMERIC_COLS if c.startswith("over4h_")]

PERIOD_RE = re.compile(r"MSitAE-([A-Za-z]+)-(\d{4})")
FILE_RE = re.compile(r"^(\d{4})-(\d{2})_monthly_ae\.csv$")
SEASONS = {12: "Winter", 1: "Winter", 2: "Winter", 3: "Spring", 4: "Spring", 5: "Spring",
           6: "Summer", 7: "Summer", 8: "Summer", 9: "Autumn", 10: "Autumn", 11: "Autumn"}


def to_number(series):
    """Turn a text column into numbers: strip commas/spaces; blanks become NaN."""
    raw = series.str.replace(",", "", regex=False).str.strip()
    return raw, pd.to_numeric(raw.where(raw.ne("")), errors="coerce")


def load_raw_files():
    """Read every raw CSV as text, keep the 22 known columns, and tag each row with its file's month."""
    files = sorted(RAW_DIR.glob("*_monthly_ae.csv"))
    if not files:
        sys.exit(f"No raw files found in {RAW_DIR}")

    frames = []
    for path in files:
        df = pd.read_csv(path, dtype=str, encoding="utf-8-sig", keep_default_na=False)
        df.columns = [c.strip() for c in df.columns]
        missing = [c for c in COLUMN_MAP if c not in df.columns]
        if missing:
            sys.exit(f"STOP: {path.name} is missing expected columns: {missing}")
        df = df[list(COLUMN_MAP)].rename(columns=COLUMN_MAP)
        m = FILE_RE.match(path.name)
        df["file_month"] = pd.Timestamp(int(m.group(1)), int(m.group(2)), 1)
        frames.append(df)

    combined = pd.concat(frames, ignore_index=True)
    print(f"Loaded {len(files)} files, {len(combined):,} rows")
    return combined


def clean(df):
    """Fix text, dates and numbers. Stop loudly on anything unexpected; log anything we fix."""
    log = []

    # 1. Trim stray spaces (e.g. "NHS ENGLAND SOUTH EAST ")
    for col in LABEL_COLS:
        df[col] = df[col].str.strip()

    # 2. Set aside NHS England's TOTAL rows. They are spelled inconsistently ("TOTAL", "Total",
    #    "TOTAl") and appear in different columns (Period in most files; Org Code/name/region in
    #    Feb 2026). Loading them would double national figures; we keep them for reconciliation.
    is_total = df[LABEL_COLS].apply(lambda col: col.str.lower().eq("total")).any(axis=1)
    totals = df[is_total].copy()
    if is_total.any():
        log.append(f"Set aside {is_total.sum()} TOTAL row(s) - used for reconciliation, not loaded")
        df = df[~is_total].copy()

    # 3. Drop rows with no organisation code
    no_code = df["org_code"].eq("")
    if no_code.any():
        log.append(f"Dropped {no_code.sum()} row(s) with no Org Code")
        df = df[~no_code].copy()

    # 4. Convert "MSitAE-SEPTEMBER-2026" -> 2026-09-01, and check it matches the file's month
    parts = df["period_raw"].str.extract(PERIOD_RE)
    df["period"] = pd.to_datetime(parts[0].str.title() + " " + parts[1], format="%B %Y", errors="coerce")
    if df["period"].isna().any():
        examples = df.loc[df["period"].isna(), "period_raw"].unique()[:3]
        sys.exit(f"STOP: could not read Period values such as {list(examples)}")
    mismatch = df["period"] != df["file_month"]
    if mismatch.any():
        sys.exit(f"STOP: {mismatch.sum()} row(s) have a Period that doesn't match their file's month")

    # 5. Convert numbers; blanks -> 0 (logged); anything non-numeric or negative stops the script
    for col in NUMERIC_COLS:
        raw, nums = to_number(df[col])
        blanks = raw.eq("")
        invalid = nums.isna() & ~blanks
        if invalid.any():
            sys.exit(f"STOP: non-numeric values in {col}, e.g. {list(raw[invalid].unique()[:5])}")
        if blanks.any():
            log.append(f"{col}: {blanks.sum()} blank value(s) set to 0")
        df[col] = nums.fillna(0).round().astype("int64")
        if (df[col] < 0).any():
            sys.exit(f"STOP: negative values found in {col}")

    # 6. One row per organisation per month
    dupes = df.duplicated(["period", "org_code"], keep="first")
    if dupes.any():
        log.append(f"Dropped {dupes.sum()} duplicate organisation-month row(s)")
        df = df[~dupes].copy()

    return df, totals, log


def reconcile_totals(df, totals):
    """Check the sum of all organisations equals NHS England's own TOTAL row, month by month."""
    print("\nReconciliation against NHS TOTAL rows:")
    if totals.empty:
        print("  No TOTAL rows found to reconcile against")
        return

    t = totals.copy()
    for col in ATT_COLS + OVER4H_COLS:
        t[col] = to_number(t[col])[1].fillna(0)

    for label, cols in [("attendances", ATT_COLS), ("over 4 hours", OVER4H_COLS)]:
        nhs = t.groupby("file_month")[cols].sum().sum(axis=1)
        ours = df.groupby("period")[cols].sum().sum(axis=1)
        comp = pd.DataFrame({"nhs_total": nhs, "our_sum": ours}).dropna()
        comp["diff"] = comp["our_sum"] - comp["nhs_total"]
        bad = comp[comp["diff"] != 0]
        status = "all match exactly" if bad.empty else f"{len(bad)} month(s) differ"
        print(f"  {label:<13}: {len(comp)} month(s) checked - {status}")
        if not bad.empty:
            print("    " + bad.head(10).to_string().replace("\n", "\n    "))

    months_without_total = sorted(set(df["period"]) - set(t["file_month"]))
    if months_without_total:
        print(f"  {len(months_without_total)} month(s) have no TOTAL row, so weren't reconciled")


def build_tables(df):
    """Split the clean data into a star schema: one fact table and two dimension tables."""
    # Fact: the numbers, keyed by month and organisation
    fact = (df[["period", "org_code"] + NUMERIC_COLS]
            .sort_values(["period", "org_code"])
            .reset_index(drop=True))

    # Organisation dimension: latest known name/region per code, plus when it appears in the data
    latest = df.sort_values("period").groupby("org_code").tail(1)
    span = df.groupby("org_code")["period"].agg(first_period="min", last_period="max")
    dim_org = (latest[["org_code", "org_name", "region"]]
               .set_index("org_code").join(span).reset_index()
               .sort_values("org_name").reset_index(drop=True))
    dim_org["region_short"] = (dim_org["region"]
                               .str.replace("NHS ENGLAND ", "", regex=False)
                               .str.title()
                               .str.replace(" And ", " and ", regex=False)
                               .str.replace(" Of ", " of ", regex=False))

    # Date dimension: one row per month, with calendar and NHS financial-year (April-March) fields
    dim_date = pd.DataFrame({"period": pd.date_range(fact["period"].min(), fact["period"].max(), freq="MS")})
    dim_date["year"] = dim_date["period"].dt.year
    dim_date["month_num"] = dim_date["period"].dt.month
    dim_date["month_name"] = dim_date["period"].dt.strftime("%B")
    dim_date["month_short"] = dim_date["period"].dt.strftime("%b")
    dim_date["month_label"] = dim_date["period"].dt.strftime("%b %Y")
    dim_date["year_month_sort"] = dim_date["year"] * 100 + dim_date["month_num"]
    dim_date["quarter"] = "Q" + dim_date["period"].dt.quarter.astype(str)
    fy_start = dim_date["year"] - (dim_date["month_num"] < 4).astype(int)
    dim_date["nhs_fy"] = fy_start.astype(str) + "-" + (fy_start + 1).astype(str).str[-2:]
    dim_date["nhs_fy_month"] = (dim_date["month_num"] - 4) % 12 + 1
    dim_date["season"] = dim_date["month_num"].map(SEASONS)

    return fact, dim_org, dim_date


def write_outputs(tables):
    """Save one Excel workbook (each sheet formatted as an Excel table) plus a CSV per table."""
    CLEAN_DIR.mkdir(parents=True, exist_ok=True)
    xlsx_path = CLEAN_DIR / "nhs_ae_model.xlsx"

    with pd.ExcelWriter(xlsx_path, engine="openpyxl") as writer:
        for name, df in tables.items():
            df.to_excel(writer, sheet_name=name, index=False)
            ws = writer.sheets[name]
            ref = f"A1:{get_column_letter(df.shape[1])}{df.shape[0] + 1}"
            table = Table(displayName=name, ref=ref)
            table.tableStyleInfo = TableStyleInfo(name="TableStyleMedium2", showRowStripes=True)
            ws.add_table(table)

    for name, df in tables.items():
        df.to_csv(CLEAN_DIR / f"{name}.csv", index=False)

    print(f"Saved {xlsx_path.relative_to(PROJECT)} and {len(tables)} CSVs")


def main():
    raw = load_raw_files()
    df, totals, log = clean(raw)
    reconcile_totals(df, totals)
    fact, dim_org, dim_date = build_tables(df)

    print()
    write_outputs({"fact_ae": fact, "dim_org": dim_org, "dim_date": dim_date})

    print("\nCleaning log:")
    for line in log or ["No fixes needed"]:
        print(f"  - {line}")

    print("\nSummary:")
    print(f"  fact_ae : {len(fact):,} rows")
    print(f"  dim_org : {len(dim_org):,} organisations")
    print(f"  dim_date: {len(dim_date)} months ({dim_date['period'].min():%b %Y} to {dim_date['period'].max():%b %Y})")
    print(f"  regions : {', '.join(sorted(dim_org['region_short'].unique()))}")

    # Sanity check to compare later with Power BI and NHS England's published figure
    latest = fact[fact["period"] == fact["period"].max()]
    att = latest[ATT_COLS].to_numpy().sum()
    over = latest[OVER4H_COLS].to_numpy().sum()
    print(f"\nSanity check - {fact['period'].max():%B %Y} national 4-hour performance (all types): {1 - over / att:.1%}")


if __name__ == "__main__":
    main()
